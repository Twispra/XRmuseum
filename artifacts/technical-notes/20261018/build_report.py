from pathlib import Path
import json
import re
import hashlib
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[3]
WORK = ROOT / 'artifacts/technical-notes/20261018'
OUT = ROOT / 'phase2小组和个人作业/COMP5424-Group2-Phase2-PrototypeDesignTechnicalNote.docx'
SOURCE_MD = ROOT / 'docs/phase2-prototype-design-and-technical-note.en.md'
DATA = json.loads((WORK/'report-content.json').read_text(encoding='utf-8'))
REFS = {r['id']: r for r in DATA['references']}
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2.5)
sec.header_distance = sec.footer_distance = Cm(1.2)

for style in doc.styles:
    if style.type in (1, 2):
        style.font.name = 'Times New Roman'
        style.font.size = Pt(12)
        style.font.color.rgb = RGBColor(0, 0, 0)
        rp = style.element.get_or_add_rPr()
        fonts = rp.find(qn('w:rFonts'))
        if fonts is None: fonts = OxmlElement('w:rFonts'); rp.insert(0, fonts)
        for key in ('ascii','hAnsi','eastAsia','cs'): fonts.set(qn('w:'+key), 'Times New Roman')
        for key in list(fonts.attrib):
            if 'Theme' in key or 'theme' in key: del fonts.attrib[key]
normal = doc.styles['Normal']
normal.paragraph_format.line_spacing = 1.04
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.widow_control = True
for key, size in [('Title',16),('Heading 1',14),('Heading 2',12)]:
    st = doc.styles[key]
    st.font.size, st.font.bold = Pt(size), True
    st.paragraph_format.space_before = Pt(5 if key!='Title' else 0)
    st.paragraph_format.space_after = Pt(5)
    st.paragraph_format.line_spacing = 1.0
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.keep_together = True
    borders = st.element.find('.//'+qn('w:pBdr'))
    if borders is not None: borders.getparent().remove(borders)
doc.core_properties.title = DATA['title']
doc.core_properties.author = 'COMP5424 Group 2'
doc.core_properties.last_modified_by = 'COMP5424 Group 2'
doc.core_properties.subject = 'Phase 2 prototype design, implementation, evidence and evaluation priorities'

def link(paragraph, text, url):
    relation = paragraph.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), relation)
    r = OxmlElement('w:r'); prop = OxmlElement('w:rPr')
    f = OxmlElement('w:rFonts')
    for key in ('ascii','hAnsi','cs','eastAsia'): f.set(qn('w:'+key),'Times New Roman')
    prop.append(f)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'),'24'); prop.append(sz)
    color = OxmlElement('w:color'); color.set(qn('w:val'),'000000'); prop.append(color)
    underline = OxmlElement('w:u'); underline.set(qn('w:val'),'single'); prop.append(underline)
    r.append(prop)
    t = OxmlElement('w:t'); t.text = text; r.append(t); h.append(r); paragraph._p.append(h)

def text(paragraph, value):
    cursor = 0
    for m in re.finditer(r'\[(\d+(?:,\d+)*)\]',value):
        paragraph.add_run(value[cursor:m.start()])
        numbers = [int(n) for n in m.group(1).split(',')]
        if len(numbers)==1 and REFS[numbers[0]].get('links'):
            link(paragraph,m.group(),REFS[numbers[0]]['links'][0][1])
        else: paragraph.add_run(m.group())
        cursor=m.end()
    paragraph.add_run(value[cursor:])

def paragraph(value, style=None):
    p=doc.add_paragraph(style=style)
    text(p,value)
    return p

def table(block):
    tbl=doc.add_table(rows=1, cols=len(block['headers']))
    tbl.alignment=WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit=False
    pr=tbl._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+side); e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side in ('top','bottom','left','right'):
        e=OxmlElement('w:'+side);e.set(qn('w:w'),'65' if side in ('top','bottom') else '80');e.set(qn('w:type'),'dxa');margins.append(e)
    pr.append(margins)
    for column,width in zip(tbl.columns,block['widths']):column.width=Cm(width)
    rows=[block['headers']]+block['rows']
    for ri,values in enumerate(rows):
        row=tbl.rows[0] if ri==0 else tbl.add_row()
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
        for ci,(cell,value) in enumerate(zip(row.cells,values)):
            cell.width=Cm(block['widths'][ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0]
            p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.space_before=Pt(0)
            p.paragraph_format.line_spacing=1.0
            p.paragraph_format.keep_with_next=False
            text(p,value)
            for r in p.runs:r.font.size=Pt(12)
            if ri==0:
                for r in p.runs:r.bold=True
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F2F2F2');cell._tc.get_or_add_tcPr().append(shade)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.paragraph_format.space_after=Pt(0)
footer.add_run('XRmuseum  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

doc.add_paragraph(DATA['title'],style='Title')
paragraph(DATA['subtitle'])
paragraph(DATA['version'])
for page_index,page in enumerate(DATA['pages']):
    heading=doc.add_paragraph(page['heading'],style='Heading 1')
    if page_index:heading.paragraph_format.page_break_before=True
    for b in page['blocks']:
        if b['type']=='p':paragraph(b['text'])
        elif b['type']=='h2':doc.add_paragraph(b['text'],style='Heading 2')
        elif b['type']=='table':table(b)
ref_heading=doc.add_paragraph('References',style='Heading 1')
ref_heading.paragraph_format.page_break_before=True
for ref in DATA['references']:
    p=paragraph('['+str(ref['id'])+'] '+ref['text'])
    p.paragraph_format.keep_together=True
    for i,(label,url) in enumerate(ref.get('links',[])):
        p.add_run(' ' if i==0 else '; ');link(p,label,url)
settings=doc.settings.element
up=OxmlElement('w:updateFields');up.set(qn('w:val'),'true');settings.append(up)
doc.save(OUT)

md=['# '+DATA['title'],'',DATA['subtitle'],'',DATA['version'],'']
for page in DATA['pages']:
    md+=['## '+page['heading'],'']
    for b in page['blocks']:
        if b['type']=='p':md +=[b['text'],'']
        elif b['type']=='h2':md+=['### '+b['text'],'']
        elif b['type']=='table':
            md+=['| '+' | '.join(b['headers'])+' |','| '+' | '.join(['---']*len(b['headers']))+' |']
            md+=['| '+' | '.join(row)+' |' for row in b['rows']]
            md+=['']
md+=['## References','']
for ref in DATA['references']:
    links='; '.join('['+label+']('+url+')' for label,url in ref.get('links',[]))
    md+=['['+str(ref['id'])+'] '+ref['text']+(' '+links if links else ''),'']
SOURCE_MD.write_text('\n'.join(md),encoding='utf-8',newline='\n')
audit={'docx':str(OUT),'docxSha256':hashlib.sha256(OUT.read_bytes()).hexdigest(),'manuscript':str(SOURCE_MD),'manuscriptSha256':hashlib.sha256(SOURCE_MD.read_bytes()).hexdigest(),'bodyWords':len(' '.join(b.get('text','')+' '+' '.join(' '.join(r) for r in b.get('rows',[])) for page in DATA['pages'] for b in page['blocks']).split()),'sourceCommit':'6088c736e89d05101e09645c1796556012a6fd5d','validationCommit':'024d03942cc93d90453583e1c7e33e4eac96ecf2','font':'Times New Roman','minimumVisibleFontPt':12,'bodyPageBudget':4,'referencePagesExcluded':True,'marginsCm':[2.5,2.5,2.5,2.5]}
(WORK/'authoring-audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(audit,ensure_ascii=False,indent=2))
